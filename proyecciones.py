"""
proyecciones.py — proyecciones de puntos (Budget / Run Rate) para comparar Real vs Proyección.

Fuente: Excel que mantiene Control de Gestión en OneDrive, uno por escenario, con
las MISMAS solapas y formato (filas = país × clave, columnas = meses abr→mar):
  Input_budget.xlsx   -> escenario 'budget'
  Input_runrate.xlsx  -> escenario 'runrate'   (opcional: si no existe, se omite)

Solapas que se leen (las demás se ignoran). Las columnas se buscan por NOMBRE
de encabezado (no por posición), los meses son las columnas cuyo encabezado es fecha:
  Cobrands_Partners  Pais, Metrica, Partner        Metrica 'Cantidad puntos' -> bloque acum_partner
                                                   Metrica 'Precio_Facturacion' (USD/pt) -> precio_partner
                                                   (Partner 'Todos' = precio de los partners del país)
  Redenciones        Pais, Metrica, Tipo_punto     puntos redimidos por tipo       -> reden
  Breakage           Pais, Metrica                 puntos vencidos                 -> breakage
  SSP_PrecioPasivo   Pais, Metrica                 SSP_Red / SSP / Precio Punto Pasivo -> ssp
  Ratio_acumulacion  LOB-Canal, Pais, Producto, Metrica, Viaje
                                                   regla de acumulación (pts / USD GB), SOLO
                                                   B2C-SITE -> ratio, clave 'Producto|Viaje'
  Penetracion_GB     LOB-Canal, Pais, Producto, Metrica, Viaje
                                                   % del GB que acumula (fracción), SOLO
                                                   LOB-Canal 'B2C-SITE' (decisión 12-oct: el GB
                                                   de Budget no viene abierto por canal) -> penet,
                                                   clave 'Producto|Domestic' / 'Producto|International'

Salida: loyalty_proyecciones.json  {meta, cols, rows} con una fila por
escenario × bloque × país × clave × mes. El cruce contra los reales lo hace
dashboard.html (selector "Comparar contra" de Métricas de Negocio) con los JSON que ya existen.

Uso:
  py proyecciones.py            # arma el JSON y muestra un resumen (no sube)
  py proyecciones.py --upload   # además lo sube a Drive (folder de loyalty)
Lo importa también loyalty_sync.py, que lo sube en cada corrida.
"""
import json
import os
import sys
from datetime import datetime, date
from pathlib import Path

_SCRIPT_DIR = Path(__file__).resolve().parent
_ONEDRIVE_DIR = (Path.home() / "OneDrive - despegar365" / "Control de Gestión - Loyalty"
                 / "Loyalty_Ecosystem")

PROYECCIONES_FILE = "loyalty_proyecciones.json"

# Por escenario: env -> OneDrive -> copia local del repo (fallback, igual que
# Input_Precios.xlsx). La máquina del Task Scheduler puede no tener OneDrive.
ESCENARIOS = {
    "budget":  [os.environ.get("BUDGET_XLSX", ""),
                str(_ONEDRIVE_DIR / "Input_budget.xlsx"),
                str(_SCRIPT_DIR / "Input_budget.xlsx")],
    "runrate": [os.environ.get("RUNRATE_XLSX", ""),
                str(_ONEDRIVE_DIR / "Input_runrate.xlsx"),
                str(_SCRIPT_DIR / "Input_runrate.xlsx")],
}

# solapa -> lista de (columnas que arman la clave, filtro de fila, bloque por fila)
def _bloque_cp(f):  # Cobrands_Partners: según Metrica
    m = (f.get("Metrica") or "").lower()
    return "precio_partner" if m.startswith("precio") else "acum_partner"

# Ratio y Penetración vienen por LOB-Canal × Viaje: solo B2C-SITE (12-oct, el GB de
# Budget no viene abierto por canal). Si la solapa no tiene LOB-Canal, pasa todo.
def _solo_site(f):
    c = (f.get("LOB-Canal") or "").upper()
    return c in ("", "B2C-SITE")

SOLAPAS = {
    "Cobrands_Partners": (["Partner"], None, _bloque_cp),
    "Redenciones":       (["Tipo_punto"], None, "reden"),
    "Breakage":          ([], None, "breakage"),
    "SSP_PrecioPasivo":  (["Metrica"], None, "ssp"),
    "Ratio_acumulacion": (["Producto", "Viaje"], _solo_site, "ratio"),
    "Penetracion_GB":    (["Producto", "Viaje"], _solo_site, "penet"),
}
# Celdas que el modelo de Planning deja sin dato: se ignoran sin warning.
_SIN_DATO = {"#missing", "#n/a", "#value!", "#div/0!", "-"}

# Nombre de país del Excel -> dataKey del dashboard (COUNTRY_DATAKEYS).
PAIS = {"argentina": "argentina", "brasil": "brasil", "brazil": "brasil",
        "mexico": "mexico", "méxico": "mexico", "colombia": "colombia",
        "peru": "peru", "perú": "peru", "ecuador": "ecuador",
        "uruguay": "uruguay", "chile": "chile"}


def _ym(v):
    if isinstance(v, (datetime, date)):
        return f"{v.year:04d}-{v.month:02d}"
    s = str(v or "").strip()
    return s[:7] if len(s) >= 7 and s[4] == "-" else None


def _read_xlsx(path, escenario, warns):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = []
    for solapa, (kcols, filtro, bloque) in SOLAPAS.items():
        if solapa not in wb.sheetnames:
            warns.append(f"{escenario}: falta la solapa '{solapa}'")
            continue
        it = wb[solapa].iter_rows(values_only=True)
        header = next(it, None) or ()
        names = {str(h).strip(): i for i, h in enumerate(header) if h is not None and not _ym(h)}
        meses = [(i, _ym(h)) for i, h in enumerate(header) if _ym(h)]
        faltan = [c for c in ["Pais"] + kcols if c not in names]
        if faltan:
            warns.append(f"{escenario}/{solapa}: faltan columnas {faltan} (encabezado: {list(names)})")
            continue
        for r in it:
            if not r:
                continue
            f = {n: (str(r[i]).strip() if i < len(r) and r[i] is not None else "") for n, i in names.items()}
            pais = PAIS.get(f["Pais"].lower())
            if pais is None:
                # 'Other Countries' y similares: no hay solapa de país en el dashboard
                continue
            if filtro and not filtro(f):
                continue
            clave = "|".join(f[c] for c in kcols)
            blq = bloque(f) if callable(bloque) else bloque
            for i, ym in meses:
                v = r[i] if i < len(r) else None
                if v is None or v == "" or str(v).strip().lower() in _SIN_DATO:
                    continue
                try:
                    v = float(v)
                except (TypeError, ValueError):
                    warns.append(f"{escenario}/{solapa}: valor no numérico {v!r} ({f['Pais']} {clave} {ym})")
                    continue
                rows.append([escenario, blq, pais, clave, ym, v])
    wb.close()
    return rows


def build_proyecciones_json():
    """Devuelve (bytes del JSON | None, resumen legible). Nunca rompe la corrida.
    Por escenario prueba cada candidato en orden (OneDrive -> copia del repo): si
    uno existe pero no se puede leer (bloqueado, sin descargar, formato roto), pasa
    al siguiente. Si no sale NINGUNA fila, devuelve None: el caller no sube nada y
    queda en Drive el loyalty_proyecciones.json anterior (no se pisa con uno vacío)."""
    rows, warns, fuentes = [], [], {}
    for esc, cands in ESCENARIOS.items():
        fuentes[esc] = None
        for path in [p for p in cands if p and Path(p).exists()]:
            try:
                esc_rows = _read_xlsx(path, esc, warns)
            except Exception as e:  # archivo abierto/bloqueado, formato roto, openpyxl faltante...
                warns.append(f"{esc}: no se pudo leer {path}: {str(e)[:150]}")
                continue
            if not esc_rows:
                warns.append(f"{esc}: {path} no trajo filas")
                continue
            rows += esc_rows
            fuentes[esc] = path
            if Path(path).parent == _SCRIPT_DIR:
                mod = datetime.fromtimestamp(Path(path).stat().st_mtime).strftime("%Y-%m-%d")
                warns.append(f"{esc}: se usó la COPIA DEL REPO ({Path(path).name}, modificada {mod}) — "
                             "puede estar desactualizada respecto de OneDrive")
            break
        if fuentes[esc] is None and esc == "budget":
            warns.append("budget: Input_budget.xlsx no encontrado o ilegible en todos los candidatos")
    if not rows:
        return None, [f"  [WARN] proyecciones: {w}" for w in warns] + [
            "  [WARN] proyecciones: 0 filas — NO se sube loyalty_proyecciones.json (queda el anterior en Drive)"]
    meta = {
        "generated_at": datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
        "fuentes": fuentes,
        "warnings": warns,
        "unidades": {"acum_partner": "puntos", "precio_partner": "USD por punto", "reden": "puntos",
                     "breakage": "puntos", "ssp": "USD por punto", "ratio": "puntos / USD GB",
                     "penet": "fracción del GB (B2C-SITE)"},
    }
    cols = ["escenario", "bloque", "pais", "clave", "ym", "valor"]
    data = json.dumps({"meta": meta, "cols": cols, "rows": rows}, ensure_ascii=False,
                      separators=(",", ":"), allow_nan=False).encode("utf-8")
    por = {}
    for r in rows:
        por[(r[0], r[1])] = por.get((r[0], r[1]), 0) + 1
    resumen = [f"  {esc:8s} {blq:13s} {n:>5} filas" for (esc, blq), n in sorted(por.items())]
    resumen += [f"  [WARN] proyecciones: {w}" for w in warns]
    return data, resumen


if __name__ == "__main__":
    data, resumen = build_proyecciones_json()
    print(f"{PROYECCIONES_FILE}: {len(data) // 1024 if data else 0} KB")
    print("\n".join(resumen))
    if data is None:
        sys.exit(1)
    if "--upload" in sys.argv:
        from google.oauth2.credentials import Credentials
        from googleapiclient.discovery import build
        from googleapiclient.http import MediaInMemoryUpload
        folder = "1yCPp6hTusYmhhb17WiB6EuhFmsx7tlxb"  # = DRIVE_FOLDER_ID de loyalty_sync.py
        svc = build("drive", "v3", credentials=Credentials.from_authorized_user_file(
            str(_SCRIPT_DIR / "token_drive.json")))
        media = MediaInMemoryUpload(data, mimetype="application/json", resumable=False)
        ex = svc.files().list(q=f"name='{PROYECCIONES_FILE}' and '{folder}' in parents and trashed=false",
                              fields="files(id)").execute().get("files", [])
        if ex:
            svc.files().update(fileId=ex[0]["id"], media_body=media).execute()
            print(f"OK Drive: actualizado ({PROYECCIONES_FILE})")
        else:
            svc.files().create(body={"name": PROYECCIONES_FILE, "parents": [folder]},
                               media_body=media, fields="id").execute()
            print(f"OK Drive: creado ({PROYECCIONES_FILE})")
