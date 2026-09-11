"""
Chequeo de reconciliacion antes de escribir_asientos.py (pedido de Rosario, 2026-08-03,
ver Bitacora seccion 42): compara lo que quedo guardado en estado_intermedio (lo que
Python calculo en la corrida de run_cierre_backup_manual.py) contra lo que HAY AHORA
en las 4 Tablas del archivo de cierre que dependen de esa corrida - Acumulacion Non
Tender (TablaANT), Redenciones (TablaRed), Redenciones SUBS (TablaRedSUBS) y
Redenciones Otros (TablaRedOtros).

Motivo: Rosario puede editar/borrar filas a mano en el archivo de cierre entre la
corrida de run_cierre_backup_manual.py y la de escribir_asientos.py (ej. sacar
transacciones puntuales de una de estas 4 hojas). Sin este chequeo, escribir_asientos.py
seguiria de largo con el estado VIEJO guardado en el pickle, ignorando en silencio
cualquier edicion manual.

Alcance del chequeo (pedido explicito de Rosario, 2026-08-03):
- SOLO fila por fila, DENTRO del `ref` de la Tabla de cada hoja - nada de lo que
  quede por fuera (ej. el bloque auxiliar $AF$2:$AG$16 de TC por transaccion en
  Redenciones/SUBS/Otros, o la tabla de referencia manual al costado de TablaANT
  que motivo el bug de la seccion 37).
- Comparacion por columna, solo de las columnas que existen en AMBOS lados (el
  DataFrame guardado y la Tabla actual). Las columnas 100% Etapa 2 (ej. "Puntos
  Valuados"/"DRO"/etc en Acumulacion Non Tender) no existen en el DataFrame
  guardado - quedan afuera de la comparacion solas, por no tener con que
  compararse. Excepcion a mano: "Reconocimiento de Ingresos Diferidos ML" en
  Redenciones/Redenciones SUBS SI existe en el DataFrame guardado, pero como
  placeholder fijo (None, pendiente de Etapa 2 - ver etapa1_redenciones.py) - si
  no se excluye, compararia None contra el valor real ya calculado en Excel y
  marcaria diferencia TODOS los meses, sin que eso signifique una edicion manual.
- Columnas numericas: tolerancia chica (misma logica que pipeline/validaciones.py
  TOLERANCIA) para no disparar falsos positivos por diferencias de redondeo entre
  el calculo de Python (pickle) y el de la formula de Excel (mismo valor, distinto
  motor de calculo).

Si hay diferencia en una hoja, pregunta por consola (una pregunta POR HOJA que
difiera, pedido explicito de Rosario) que hacer: seguir con lo guardado (ignorar,
fue una edicion de prueba/error) o usar la version nueva de la Tabla (la edicion
fue intencional). Devuelve siempre un DataFrame con las mismas columnas que el
guardado (si se elige "nuevo", se agregan como None las columnas Etapa-2-pendientes
que la Tabla no tiene, para no romper el resto del flujo que las espera).
"""

import re

import pandas as pd
from openpyxl import load_workbook

from pipeline_backup.escritura_cierre_manual import ruta_cierre_mensual

TOLERANCIA = 0.01

NOMBRES_TABLA = {
    "Acumulacion Non Tender": "TablaANT",
    "Redenciones": "TablaRed",
    "Redenciones SUBS": "TablaRedSUBS",
    "Redenciones Otros": "TablaRedOtros",
}

# Columnas que existen en el DataFrame guardado pero son placeholders fijos
# (None, pendientes de Etapa 2) - no comparables contra el valor real de la Tabla.
COLUMNAS_EXCLUIDAS = {
    "Redenciones": {"Reconocimiento de Ingresos Diferidos ML"},
    "Redenciones SUBS": {"Reconocimiento de Ingresos Diferidos ML"},
}


_WB_CACHE = None


def _workbook():
    """Abre el archivo de cierre UNA sola vez (cacheado a nivel modulo) y lo
    reusa para las 4 hojas. Tiene que ser modo NORMAL (no read_only): `ws.tables`
    no esta disponible en read_only (confirmado - ReadOnlyWorksheet no tiene ese
    atributo), y necesitamos el `ref` de la Tabla para saber donde limitarnos.
    Mismo costo ya aceptado en escritura_cierre_manual.pegar_todo_manual (~2.1 GB /
    30s con el archivo unificado sin pivots, ver Bitacora seccion 33) - pagarlo una
    sola vez para las 4 hojas en vez de 4 veces."""
    global _WB_CACHE
    if _WB_CACHE is None:
        _WB_CACHE = load_workbook(ruta_cierre_mensual(), data_only=True)
    return _WB_CACHE


def cerrar():
    """Libera el workbook cacheado - llamar al terminar de reconciliar las 4 hojas."""
    global _WB_CACHE
    if _WB_CACHE is not None:
        _WB_CACHE.close()
        _WB_CACHE = None


def _leer_tabla(nombre_hoja: str) -> pd.DataFrame:
    """Lee TODA la Tabla (headers + filas), tal cual esta AHORA en el archivo de
    cierre, limitado al `ref` declarado de esa Tabla - nada de por fuera (ver
    docstring del modulo)."""
    ws = _workbook()[nombre_hoja]
    tabla = ws.tables[NOMBRES_TABLA[nombre_hoja]]
    col_inicio_ref, ref_fin = tabla.ref.split(":")
    letra_col_inicio = re.match(r"[A-Za-z]+", col_inicio_ref).group()
    letra_col_fin = re.match(r"[A-Za-z]+", ref_fin).group()
    fila_fin = int(re.search(r"\d+", ref_fin).group())

    col_inicio_idx = ws[letra_col_inicio + "1"].column
    col_fin_idx = ws[letra_col_fin + "1"].column

    filas = ws.iter_rows(
        min_row=1, max_row=fila_fin, min_col=col_inicio_idx, max_col=col_fin_idx,
        values_only=True,
    )
    header = next(filas)
    datos = {h: [] for h in header if h is not None}
    for fila in filas:
        for h, v in zip(header, fila):
            if h is not None:
                datos[h].append(v)
    return pd.DataFrame(datos)


def _columna_igual(a: pd.Series, b: pd.Series) -> bool:
    """Compara 2 columnas fila a fila. Numericas: con tolerancia (redondeo Python
    vs formula de Excel). El resto: exacto, tratando None/NaN como iguales entre si."""
    a_na = a.isna()
    b_na = b.isna()
    if not (a_na == b_na).all():
        return False
    a_validos = a[~a_na]
    b_validos = b[~b_na]
    if pd.api.types.is_numeric_dtype(a_validos) and pd.api.types.is_numeric_dtype(b_validos):
        return bool((a_validos.astype(float) - b_validos.astype(float)).abs().le(TOLERANCIA).all())
    return bool((a_validos.astype(str) == b_validos.astype(str)).all())


def _comparar(df_guardado: pd.DataFrame, df_archivo: pd.DataFrame, nombre_hoja: str) -> bool:
    """True si coinciden (en las columnas comparables) - False si difieren."""
    if len(df_guardado) != len(df_archivo):
        return False
    excluidas = COLUMNAS_EXCLUIDAS.get(nombre_hoja, set())
    columnas_comunes = [
        c for c in df_guardado.columns if c in df_archivo.columns and c not in excluidas
    ]
    g = df_guardado.reset_index(drop=True)
    f = df_archivo.reset_index(drop=True)
    return all(_columna_igual(g[c], f[c]) for c in columnas_comunes)


def _preguntar_como_seguir(nombre_hoja: str) -> str:
    print(f"\n  DIFERENCIA detectada en '{nombre_hoja}': lo guardado de la corrida de")
    print("  run_cierre_backup_manual.py no coincide con lo que hay AHORA en esa Tabla")
    print("  del archivo de cierre (¿se edito/borro algo a mano?).")
    while True:
        respuesta = input(
            "  [1] Fue un error / lo voy a re-chequear despues - pegar la version GUARDADA de la bajada\n"
            "  [2] El cambio es intencional - usar la version NUEVA del archivo de cierre\n"
            "  Elegi 1 o 2: "
        ).strip()
        if respuesta in ("1", "2"):
            return "guardado" if respuesta == "1" else "nuevo"
        print("  Respuesta invalida - escribi 1 o 2.")


def reconciliar(nombre_hoja: str, df_guardado: pd.DataFrame) -> pd.DataFrame:
    """Compara df_guardado contra la Tabla de `nombre_hoja` en el archivo de cierre
    tal cual esta ahora. Si coinciden, devuelve df_guardado sin cambios. Si difieren,
    pregunta por consola y devuelve la version elegida - siempre con las mismas
    columnas que df_guardado (si se elige "nuevo", las columnas Etapa-2-pendientes
    que la Tabla no tiene se agregan como None)."""
    df_archivo = _leer_tabla(nombre_hoja)

    if _comparar(df_guardado, df_archivo, nombre_hoja):
        return df_guardado

    eleccion = _preguntar_como_seguir(nombre_hoja)
    if eleccion == "guardado":
        print(f"  -> Sigo con la version guardada de la corrida para '{nombre_hoja}'.")
        return df_guardado

    print(f"  -> Uso la version nueva del archivo de cierre para '{nombre_hoja}'.")
    for columna in df_guardado.columns:
        if columna not in df_archivo.columns:
            df_archivo[columna] = None
    return df_archivo[df_guardado.columns]
