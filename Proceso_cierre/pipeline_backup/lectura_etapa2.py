"""
Lee del archivo de cierre (unificado desde 2026-07-27, ver Bitacora seccion 33), YA
RECALCULADO en Excel de verdad tras el refresco manual (Ctrl+Alt+F9), las columnas
que hasta ahora quedaban "pendientes de Etapa 2" - ver Bitacora seccion 31 para la
cadena de dependencia completa (Breakage Esperado -> SSP Facturacion ->
Cobrand/Partners/Puntos Valuados etc.) y la decision de por que el recalculo tiene
que ser manual (openpyxl no recalcula formulas).

Solo LEE este archivo, nunca escribe - siempre con read_only=True, data_only=True
(mas liviano todavia que el modo escritura, que ya de por si es seguro ahora que el
archivo no tiene pivots - ver seccion 33).
"""

import openpyxl
import pandas as pd

from pipeline.config import cargar_config, ruta_cierre

COLUMNAS_ACUMULACION_NON_TENDER = ["Puntos Valuados", "DRO", "DRO - up fronts", "DRO -Fee + Descuentos"]
COLUMNA_RECONOCIMIENTO_ML = "Reconocimiento de Ingresos Diferidos ML"

# Bloque manual de pivot "Breakage por pais/Entidad Legal" al costado de TablaBr, en
# la hoja "Breakage" del archivo de cierre - headers en fila 3 (no fila 1, por eso
# _leer_columnas_por_header necesita fila_header/min_col/max_col parametrizables,
# ver seccion 44 de la Bitacora). Limitado a columnas K:R para no mezclarse con la
# tabla principal (A:F) ni con las columnas de ayuda/lookup mas a la derecha (S en
# adelante) - mismo motivo que el bug de la seccion 37 (Acumulacion Non Tender).
COLUMNAS_BREAKAGE_PAIS_EL = [
    "country_code", "Entidad Legal", "channel_condition", "produto_original",
    "Producto", "Suma de DRO", "% DRO", "Breakage por país y LE (USD)",
]
BREAKAGE_FILA_HEADER = 3
BREAKAGE_MIN_COL = 11  # K
BREAKAGE_MAX_COL = 18  # R


def _ruta_archivo_cierre() -> str:
    cfg = cargar_config()
    return ruta_cierre(cfg["backup_manual"]["archivo_cierre"])


def _leer_columnas_por_header(ruta: str, nombre_hoja: str, columnas: list,
                               fila_header: int = 1, min_col: int = None, max_col: int = None) -> pd.DataFrame:
    """Lee columnas por nombre de header (fila `fila_header`), en el mismo orden de
    filas en que estan en el archivo (fila `fila_header + 1` en adelante) - se corta
    en la primera fila completamente vacia (mirando solo min_col:max_col, si se
    pasan - default: toda la fila, comportamiento historico sin cambios)."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        ws = wb[nombre_hoja]
        filas = ws.iter_rows(min_row=fila_header, min_col=min_col, max_col=max_col, values_only=True)
        header = next(filas)
        idx = {h: i for i, h in enumerate(header) if h in columnas}
        faltantes = set(columnas) - set(idx)
        if faltantes:
            raise KeyError(
                f"No se encontraron los headers {sorted(faltantes)} en '{nombre_hoja}' "
                f"de {ruta} - revisar si cambio el archivo antes de seguir."
            )
        datos = {c: [] for c in columnas}
        for fila in filas:
            if all(v is None for v in fila):
                break
            for c in columnas:
                datos[c].append(fila[idx[c]])
        return pd.DataFrame(datos)
    finally:
        wb.close()


def leer_columnas_generacion() -> pd.DataFrame:
    """Puntos Valuados / DRO / DRO - up fronts / DRO -Fee + Descuentos, ya
    recalculadas en 'Acumulacion Non Tender'."""
    return _leer_columnas_por_header(_ruta_archivo_cierre(), "Acumulacion Non Tender", COLUMNAS_ACUMULACION_NON_TENDER)


def leer_reconocimiento_ml(nombre_hoja: str) -> pd.Series:
    """Reconocimiento de Ingresos Diferidos ML, ya recalculado, para 'Redenciones'
    o 'Redenciones SUBS'."""
    df = _leer_columnas_por_header(_ruta_archivo_cierre(), nombre_hoja, [COLUMNA_RECONOCIMIENTO_ML])
    return df[COLUMNA_RECONOCIMIENTO_ML]


def leer_breakage_pais_el() -> pd.DataFrame:
    """Breakage por pais/Entidad Legal (columnas K:R de 'Breakage' en el archivo de
    cierre, ya recalculado) - se pega a valor en Asientos!Breakage columnas A:H (ver
    escritura_asientos.pegar_breakage, y Proceso Cierre.pdf: 'de la K a la R de
    Breakage en el cierre van a valor en Breakage de Asientos, de la A a la H')."""
    return _leer_columnas_por_header(
        _ruta_archivo_cierre(), "Breakage", COLUMNAS_BREAKAGE_PAIS_EL,
        fila_header=BREAKAGE_FILA_HEADER, min_col=BREAKAGE_MIN_COL, max_col=BREAKAGE_MAX_COL,
    )


def leer_accounting() -> list:
    """Grilla de valores (lista de listas, una por fila) de la solapa 'Accounting',
    ya recalculada - se pega tal cual a valor en Asientos!Accounting (decision de
    Rosario, ver Bitacora seccion 31: esa hoja es una copia exacta de esta)."""
    wb = openpyxl.load_workbook(_ruta_archivo_cierre(), read_only=True, data_only=True)
    try:
        ws = wb["Accounting"]
        return [list(fila) for fila in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
