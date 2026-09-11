"""
Escribe en el archivo de cierre unificado (`backup_manual.archivo_cierre`, ej.
"Cierre 2026 07.xlsx") los pasos 1 y 2 de Proceso Cierre.pdf, para que Rosario pueda
seguir la Etapa 2 a mano (o dejar que recalculen las formulas de Etapa 2 ya
existentes - ver Bitacora seccion 31/33) usando ese archivo como punto de partida.

DIFERENCIA CLAVE con pipeline/escritura_asientos.py: ahi Python CALCULA las columnas
derivadas (Revenue, Entidad Legal, etc.) y las pega A VALOR. Aca NO - esas columnas son
FORMULAS EN VIVO que ya viven en el archivo de cierre (confirmado celda a celda,
2026-07-22: usan una mezcla de referencias de tabla estructurada `TablaXXX[[#This Row],
[columna]]` y referencias de celda relativa simple `T2`/`U2`/etc.). Por eso este modulo
recibe el DataFrame ya FILTRADO por point_type pero SIN CALCULAR (usa
etapa1_acumulaciones.filtrar_acumulacion_non_tender() / etapa1_redenciones.filtrar_*(),
no procesar()/calcular_columnas_derivadas()) - solo se pegan los campos crudos, y las
columnas de formula se replican tal cual estaban (mismo mecanismo que escritura_asientos.py:
capturar plantilla de la fila 2 ANTES de borrar, despues reemplazar el mes).

Igual que Asientos Cierre Loyalty.xlsx (ver Bitacora seccion 15/16), el archivo de
cierre es un archivo por mes (se duplica el mes anterior a mano, renombrando al mes
nuevo) - NO acumula historico entre corridas. Cada corrida REEMPLAZA los datos de
cada hoja con los del mes nuevo.

HISTORIA (2026-07-27, ver Bitacora seccion 33): hasta ahora este archivo estaba
separado en 2 (Bases liviano + Pesado con Etapa 2 y pivots) porque abrir el
combinado con openpyxl en modo escritura llegaba a 8+ GB de RAM (seccion 23). Rosario
saco TODOS los pivots del archivo (tanto de este como de Asientos) - confirmado que
el archivo unificado sin pivots abre en modo escritura con ~2.1 GB / 30s, seguro. Por
eso se volvio a un solo archivo: este modulo ahora pega TAMBIEN la solapa
"Diccionario" (antes en pipeline_backup/escritura_diccionario.py, con cirugia de XML
porque el archivo pesado no era seguro de abrir - eso ya no hace falta).

Mapeo real de las hojas que este modulo toca (confirmado 2026-07-22/27):

- "Query Acumulaciones" (Tabla `TablaQAc`) y "Query Redenciones" (Tabla `TablaQRed`):
  volcado PLANO de la bajada cruda, sin ningun filtro ni formula - los headers coinciden
  1 a 1 con los nombres de columna que ya trae la query (sin alias). Se pega el DataFrame
  crudo completo (sin filtrar por point_type).
  ⚠️ "Query Acumulaciones" puede tener remanente de meses anteriores por fuera del `ref`
  de `TablaQAc` - este modulo lo resuelve solo: la deteccion de "ultima fila a borrar" no
  confia en el `ref` de la Tabla, escanea la hoja entera por la columna A.

- "Breakage" (Tabla `TablaBr`): volcado de Puntos Expirados, con UNA columna de formula
  (F, "Aplica" = VLOOKUP contra Diccionario). El header renombra `points_type_id` a
  "Point Type" (unico alias en esta hoja).

- "Acumulacion Non Tender" (Tabla `TablaANT`): recibe el filtrado de Acumulaciones (SIN
  calcular). Raw+pass-through A-S (19 campos, headers = nombres nativos de la query, sin
  alias). T en adelante (Revenue, Puntos Valuados, DRO, DRO-up fronts, DRO-Fee+Descuentos,
  Filtro P&L, Entidad Legal, Producto, Points ABS, Producto homologado, Pts>0, GB pts>0)
  son formulas que se replican de la plantilla - Python nunca calcula estos valores aca.

- "Redenciones" / "Redenciones SUBS" (Tablas `TablaRed` / `TablaRedSUBS`): raw+pass-through
  A-T (20 campos, sin alias - los headers de esta hoja usan los nombres nativos de la
  query). U en adelante (Reconocimiento ML, Concat Entidad Legal, Producto, Reduccion del
  pasivo, DRI) son formulas replicadas.

- "Redenciones Otros" (Tabla `TablaRedOtros`): mismo patron A-T, PERO con 3 alias (igual
  que se encontro en Asientos): `country_code`->"country", `country`->"pais",
  `payment_type`->"tipopago". U en adelante (Reconocimiento ML, Concat Entidad Legal,
  Producto, Reconocimiento USD, Columna1) formulas replicadas.

- "Diccionario" (SIN Tabla, sin Pivot Cache que la use - confirmado): id/code/tratamiento,
  resultado de `Queries Data/Diccionario Puntos Excl.sql` - se pisa entera.

IMPORTANTE (ver Bitacora seccion 11): antes de bajar los datos, el orquestador
(run_cierre_backup_manual.py) ya corrio los 2 guardrails (diccionario_puntos,
validacion_queries) - la clasificacion de "Redenciones Otros" por complemento depende de
eso.
"""

import os
import re
import shutil
from datetime import datetime

import pandas as pd
from openpyxl.utils import get_column_letter

from pipeline.config import cargar_config, ruta_absoluta, ruta_cierre
from pipeline.conexion import fetch
from pipeline.escritura_asientos import (
    _ajustar_filas_formula, _set_cell, _tablas_reales, _TABLE_REF_RE, FORMATOS_FIJOS,
)


FILA_HEADER = 1
FILA_TEMPLATE = 2  # primera fila de datos en este archivo (a diferencia de Asientos, que usa fila 3)

COLUMNAS_DICCIONARIO = ["id", "code", "tratamiento"]


def _header_a_columna(ws, fila_header: int = FILA_HEADER, min_col: int = None, max_col: int = None) -> dict:
    """Mapea nombre de header -> columna, escaneando SOLO min_col:max_col (por defecto
    toda la fila, para las hojas sin Tabla como Diccionario).

    2026-07-28 (ver Bitacora seccion 37) - BUG REAL encontrado y corregido: sin limite de
    columnas, esto escaneaba la fila 1 ENTERA. En 'Acumulacion Non Tender', que comparte
    la hoja con una tabla de referencia manual mas a la derecha de TablaANT (fuera de su
    `ref`, A1:AE...) que por coincidencia tambien tiene columnas llamadas "country_code" y
    "country", el diccionario armado por este escaneo terminaba quedandose con la ULTIMA
    coincidencia (la de la derecha, ajena a TablaANT) en vez de la columna real de la
    Tabla - todos los valores de country_code/country se pegaban en la tabla de
    referencia ajena, y la columna real de TablaANT quedaba vacia. Mismo patron que el
    bug de `ultima_col_idx` (ver docstring de `_pegar_hoja`): nunca escanear mas alla de
    las columnas propias de la Tabla."""
    mapa = {}
    for cell in ws[fila_header][slice(min_col - 1 if min_col else None, max_col)]:
        if cell.value is not None:
            mapa[str(cell.value).strip()] = cell.column
    return mapa


def _ultima_fila_real(ws, columna: int = 1) -> int:
    """Ultima fila con dato real en `columna`, escaneando la hoja entera - NO confia en
    el `ref` de la Tabla (hay hojas, ej. 'Query Acumulaciones', donde el ref quedo
    desactualizado y hay datos reales por fuera de el)."""
    ultima = FILA_HEADER
    for fila in ws.iter_rows(min_row=FILA_HEADER + 1, min_col=columna, max_col=columna):
        if fila[0].value is not None:
            ultima = fila[0].row
    return ultima


MAPEO_QUERY_ACUMULACIONES = [
    ("processing_date", "processing_date"), ("country_code", "country_code"),
    ("partner", "partner"), ("points", "points"), ("point_type", "point_type"),
    ("business", "business"), ("product", "product"), ("country", "country"),
    ("payment_type", "payment_type"), ("channel_condition", "channel_condition"),
    ("trip_type", "trip_type"), ("produto_original", "produto_original"),
    ("produto", "produto"), ("comision", "comision"), ("fee", "fee"),
    ("descuentos", "descuentos"), ("gb_basebi", "gb_basebi"), ("gb_basebi_2", "gb_basebi_2"),
    ("pct_pagado_con_puntos", "pct_pagado_con_puntos"), ("concatenado", "concatenado"),
    ("legal_entity", "legal_entity"), ("descuento_consumo_puntos_usd", "descuento_consumo_puntos_usd"),
]

MAPEO_QUERY_REDENCIONES = [
    ("processing_date", "processing_date"), ("product", "product"),
    ("country_code", "country_code"), ("partner", "partner"), ("points", "points"),
    ("point_type", "point_type"), ("points_distribuidos", "points_distribuidos"),
    ("business", "business"), ("country", "country"), ("produto_original", "produto_original"),
    ("channel_condition", "channel_condition"), ("payment_type", "payment_type"),
    ("produto_agrupado", "produto_agrupado"), ("comision", "comision"), ("fee", "fee"),
    ("descuentos", "descuentos"), ("gb_basebi", "gb_basebi"), ("gb_basebi_2", "gb_basebi_2"),
    ("descuento_consumo_puntos_usd", "descuento_consumo_puntos_usd"),
    ("legal_entity", "legal_entity"), ("concatenado", "concatenado"), ("moneda_local", "moneda_local"),
]

MAPEO_BREAKAGE = [
    ("Expiration_Date", "Expiration_Date"), ("country", "country"), ("tier", "tier"),
    ("Point Type", "points_type_id"), ("Points", "Points"),
]

MAPEO_ACUMULACION_NON_TENDER = [
    ("processing_date", "processing_date"), ("country_code", "country_code"),
    ("partner", "partner"), ("points", "points"), ("point_type", "point_type"),
    ("business", "business"), ("product", "product"), ("country", "country"),
    ("payment_type", "payment_type"), ("channel_condition", "channel_condition"),
    ("trip_type", "trip_type"), ("produto_original", "produto_original"),
    ("produto", "produto"), ("comision", "comision"), ("fee", "fee"),
    ("descuentos", "descuentos"), ("gb_basebi", "gb_basebi"), ("gb_basebi_2", "gb_basebi_2"),
    ("pct_pagado_con_puntos", "pct_pagado_con_puntos"),
    # T en adelante (Revenue, Puntos Valuados, DRO, ...): formulas, no se pegan - ver plantilla.
]


def _mapeo_redenciones_familia(alias_country_code="country_code", alias_country="country",
                                alias_payment_type="payment_type"):
    return [
        ("processing_date", "processing_date"), ("product", "product"),
        (alias_country_code, "country_code"), ("partner", "partner"), ("points", "points"),
        ("point_type", "point_type"), ("points_distribuidos", "points_distribuidos"),
        ("business", "business"), (alias_country, "country"),
        ("produto_original", "produto_original"), ("channel_condition", "channel_condition"),
        (alias_payment_type, "payment_type"), ("produto_agrupado", "produto_agrupado"),
        ("comision", "comision"), ("fee", "fee"), ("descuentos", "descuentos"),
        ("gb_basebi", "gb_basebi"), ("gb_basebi_2", "gb_basebi_2"),
        ("descuento_consumo_puntos_usd", "descuento_consumo_puntos_usd"),
        ("legal_entity", "legal_entity"),
        # U en adelante (Reconocimiento ML, Concat Entidad Legal, ...): formulas, no se pegan.
    ]


MAPEO_REDENCIONES = _mapeo_redenciones_familia()
MAPEO_REDENCIONES_SUBS = _mapeo_redenciones_familia()
MAPEO_REDENCIONES_OTROS = _mapeo_redenciones_familia(
    alias_country_code="country", alias_country="pais", alias_payment_type="tipopago"
)

HOJAS_RAW = {
    # Sin filtro, sin formulas - volcado completo de la bajada cruda.
    "Query Acumulaciones": {"tabla": "TablaQAc", "mapeo": MAPEO_QUERY_ACUMULACIONES, "primera_col_formula": None},
    "Query Redenciones": {"tabla": "TablaQRed", "mapeo": MAPEO_QUERY_REDENCIONES, "primera_col_formula": None},
}

HOJAS_FILTRADAS = {
    "Breakage": {"tabla": "TablaBr", "mapeo": MAPEO_BREAKAGE, "primera_col_formula": "F"},
    "Acumulacion Non Tender": {"tabla": "TablaANT", "mapeo": MAPEO_ACUMULACION_NON_TENDER, "primera_col_formula": "T"},
    "Redenciones": {"tabla": "TablaRed", "mapeo": MAPEO_REDENCIONES, "primera_col_formula": "U"},
    "Redenciones SUBS": {"tabla": "TablaRedSUBS", "mapeo": MAPEO_REDENCIONES_SUBS, "primera_col_formula": "U"},
    "Redenciones Otros": {"tabla": "TablaRedOtros", "mapeo": MAPEO_REDENCIONES_OTROS, "primera_col_formula": "U"},
}


def _pegar_hoja(ws, config_hoja: dict, df: pd.DataFrame) -> int:
    """Reemplaza el contenido de una hoja (raw, o raw+formulas-replicadas segun
    config_hoja['primera_col_formula']) con los datos de df. Devuelve filas escritas.

    2026-07-28 (ver Bitacora seccion 34) - BUG REAL encontrado y corregido: `ultima_col_idx`
    se calculaba escaneando TODA la fila 1 de la hoja buscando cualquier header
    (`max(encabezados.values())`), sin limitarse a las columnas propias de la Tabla de esta
    hoja. En 'Breakage', que comparte la hoja con 'Tabla10' (tabla manual de Rosario en
    K3:S821), una nota suelta en K1 ("Actualizar colores...", NO un header real - el header
    real de Tabla10 esta en la fila 3) se contaba como si fuera un header mas -> el codigo
    extendia la "plantilla" de formulas replicadas hasta la columna K, pisando con `None` el
    inicio de los datos manuales de Tabla10 en las filas que se solapan con las filas nuevas
    de TablaBr. Corregido: `ultima_col_idx` ahora se deriva del `ref` DECLARADO de la Tabla
    (`letra_col_fin`, ya usado para las filas) - nunca de un escaneo ciego de la fila 1. Esto
    es ademas mas robusto en general (no depende de que la fila 1 este "limpia")."""
    tabla = ws.tables[config_hoja["tabla"]]
    col_inicio_ref, ref_fin = tabla.ref.split(":")
    letra_col_inicio = re.match(r"[A-Za-z]+", col_inicio_ref).group()
    letra_col_fin = re.match(r"[A-Za-z]+", ref_fin).group()

    col_inicio_idx = ws[letra_col_inicio + "1"].column
    ultima_col_idx = ws[letra_col_fin + "1"].column
    encabezados = _header_a_columna(ws, min_col=col_inicio_idx, max_col=ultima_col_idx)

    # Detectar la ultima fila real ESCANEANDO, no confiando en el ref de la Tabla
    # (ver docstring del modulo - 'Query Acumulaciones' tiene remanente huerfano fuera
    # del ref declarado).
    fila_vieja_max = max(
        int(re.search(r"\d+", ref_fin).group()),
        _ultima_fila_real(ws),
    )

    plantilla = {}
    if config_hoja["primera_col_formula"]:
        primera_col_formula_idx = ws[config_hoja["primera_col_formula"] + "1"].column
        for col_idx in range(primera_col_formula_idx, ultima_col_idx + 1):
            plantilla[col_idx] = ws.cell(row=FILA_TEMPLATE, column=col_idx).value

    df = df.reset_index(drop=True) if df is not None else pd.DataFrame()
    filas_nuevas = len(df)
    fila_nueva_max = FILA_HEADER + filas_nuevas

    for i, fila in df.iterrows():
        fila_excel = FILA_TEMPLATE + i

        for header_hoja, columna_df in config_hoja["mapeo"]:
            col_idx = encabezados.get(header_hoja)
            if col_idx is None:
                raise KeyError(
                    f"No se encontro el header '{header_hoja}' en la hoja del archivo "
                    "de cierre - revisar si cambio el archivo antes de seguir."
                )
            valor = fila[columna_df] if columna_df in df.columns else None
            if pd.isna(valor):
                valor = None
            _set_cell(ws, fila_excel, col_idx, valor, FORMATOS_FIJOS.get(columna_df))

        for col_idx, valor_plantilla in plantilla.items():
            if valor_plantilla is None:
                nuevo_valor = None
            elif isinstance(valor_plantilla, str) and valor_plantilla.startswith("="):
                nuevo_valor = _ajustar_filas_formula(valor_plantilla, FILA_TEMPLATE, fila_excel)
            else:
                nuevo_valor = valor_plantilla
            _set_cell(ws, fila_excel, col_idx, nuevo_valor)

    if fila_vieja_max > fila_nueva_max:
        for fila_excel in range(fila_nueva_max + 1, fila_vieja_max + 1):
            for col_idx in range(1, ultima_col_idx + 1):
                _set_cell(ws, fila_excel, col_idx, None)

    tabla.ref = f"{col_inicio_ref}:{letra_col_fin}{fila_nueva_max}"
    return filas_nuevas


def pegar_query_acumulaciones(wb, df_acumulaciones: pd.DataFrame) -> int:
    return _pegar_hoja(wb["Query Acumulaciones"], HOJAS_RAW["Query Acumulaciones"], df_acumulaciones)


def pegar_query_redenciones(wb, df_redenciones: pd.DataFrame) -> int:
    return _pegar_hoja(wb["Query Redenciones"], HOJAS_RAW["Query Redenciones"], df_redenciones)


def pegar_breakage(wb, df_puntos_expirados: pd.DataFrame) -> int:
    return _pegar_hoja(wb["Breakage"], HOJAS_FILTRADAS["Breakage"], df_puntos_expirados)


def pegar_acumulacion_non_tender(wb, df_acumulaciones: pd.DataFrame, point_types: list) -> int:
    from pipeline.etapa1_acumulaciones import filtrar_acumulacion_non_tender
    df_filtrado = filtrar_acumulacion_non_tender(df_acumulaciones, point_types)
    return _pegar_hoja(wb["Acumulacion Non Tender"], HOJAS_FILTRADAS["Acumulacion Non Tender"], df_filtrado)


def pegar_redenciones_filtradas(wb, df_redenciones: pd.DataFrame,
                                 point_types_general: list, point_types_subs: list) -> dict:
    from pipeline.etapa1_redenciones import (
        filtrar_redenciones_general, filtrar_redenciones_subs, filtrar_redenciones_otros,
    )
    df_general = filtrar_redenciones_general(df_redenciones, point_types_general)
    df_subs = filtrar_redenciones_subs(df_redenciones, point_types_subs)
    df_otros = filtrar_redenciones_otros(df_redenciones, point_types_general, point_types_subs)
    return {
        "Redenciones": _pegar_hoja(wb["Redenciones"], HOJAS_FILTRADAS["Redenciones"], df_general),
        "Redenciones SUBS": _pegar_hoja(wb["Redenciones SUBS"], HOJAS_FILTRADAS["Redenciones SUBS"], df_subs),
        "Redenciones Otros": _pegar_hoja(wb["Redenciones Otros"], HOJAS_FILTRADAS["Redenciones Otros"], df_otros),
    }


def _ruta_query_diccionario_excl() -> str:
    cfg = cargar_config()
    return ruta_absoluta("queries", cfg["queries"]["data"]["diccionario_excl"])


def bajar_diccionario_excl() -> pd.DataFrame:
    with open(_ruta_query_diccionario_excl(), encoding="utf-8") as f:
        query = f.read()
    return fetch(query, "Diccionario Puntos Excl")


def pegar_diccionario(wb, df: pd.DataFrame) -> int:
    """Pisa la hoja 'Diccionario' (sin Tabla, sin Pivot Cache que la use - confirmado
    2026-07-24/27) con id/code/tratamiento - se borra lo que hubiera antes.

    2026-07-27: antes esto vivia en pipeline_backup/escritura_diccionario.py y hacia
    cirugia de XML porque el archivo pesado no era seguro de abrir con openpyxl. Con
    el archivo unificado sin pivots eso ya no hace falta - se pega como cualquier
    otra hoja, en la misma sesion de wb que el resto (ver Bitacora seccion 33)."""
    ws = wb["Diccionario"]
    for rango in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rango))

    # Ojo: el esquema viejo (antes de la seccion 27) tenia 4 columnas (A-D). Hay que
    # limpiar cualquier columna vieja de mas, no solo filas de mas en las 3 columnas
    # actuales - si no, queda un remanente (ej. la vieja columna D) sin borrar.
    # 2026-07-28: se usa `ws.max_row`/`ws.max_column` (no `_ultima_fila_real`, que solo
    # mira VALORES) porque esta hoja tenia filas sin valor pero con formato/estilo
    # residual (ver Bitacora seccion 35) - esas filas igual cuentan para la dimension
    # real de la hoja y hay que borrarlas tambien.
    fila_vieja_max = ws.max_row
    col_vieja_max = ws.max_column

    for i, columna in enumerate(COLUMNAS_DICCIONARIO, start=1):
        _set_cell(ws, FILA_HEADER, i, columna)

    df = df.reset_index(drop=True)
    filas_nuevas = len(df)
    for idx, fila in df.iterrows():
        fila_excel = FILA_TEMPLATE + idx
        for i, columna in enumerate(COLUMNAS_DICCIONARIO, start=1):
            valor = fila[columna] if columna in df.columns else None
            if pd.isna(valor):
                valor = None
            _set_cell(ws, fila_excel, i, valor)

    fila_nueva_max = FILA_HEADER + filas_nuevas

    # 2026-07-28 (ver Bitacora seccion 35) - BUG REAL encontrado y corregido: esta hoja
    # no tiene Tabla, asi que el remanente viejo (columna D del esquema pre-seccion-27,
    # filas de mas de una corrida anterior con mas datos) no se limpiaba de verdad -
    # antes solo se ponia en None, pero la celda seguia EXISTIENDO (con su estilo), asi
    # que la `dimension` de la hoja y el `AutoFilter` propio de la hoja (`ws.auto_filter`,
    # que quedaba apuntando al rango VIEJO ej. "A1:D120") nunca se achicaban - un
    # AutoFilter/FilterDatabase apuntando a una columna sin header ni datos es lo que
    # Excel detectaba como "problema con el contenido" al abrir el archivo. Corregido:
    # se BORRAN de verdad (no solo se vacian) las columnas/filas de mas, y se re-acota
    # el AutoFilter al rango real.
    if col_vieja_max > len(COLUMNAS_DICCIONARIO):
        ws.delete_cols(len(COLUMNAS_DICCIONARIO) + 1, col_vieja_max - len(COLUMNAS_DICCIONARIO))
    if fila_vieja_max > fila_nueva_max:
        ws.delete_rows(fila_nueva_max + 1, fila_vieja_max - fila_nueva_max)

    letra_col_fin = get_column_letter(len(COLUMNAS_DICCIONARIO))
    ws.auto_filter.ref = f"A{FILA_HEADER}:{letra_col_fin}{fila_nueva_max}"

    return filas_nuevas


def ruta_cierre_mensual() -> str:
    """Ruta del archivo de cierre unificado - el UNICO que este modulo escribe. Vive
    DENTRO de rutas.carpeta_cierre (una copia por mes, igual que Asientos Cierre
    Loyalty.xlsx - ver config.yaml -> backup_manual.archivo_cierre y Bitacora
    secciones 23 y 33)."""
    cfg = cargar_config()
    return ruta_cierre(cfg["backup_manual"]["archivo_cierre"])


def hacer_backup(ruta_archivo: str) -> str:
    """Backup con timestamp antes de sobreescribir - dentro de la carpeta del mes,
    separada de la de Asientos (ver config.yaml -> backup_manual.carpeta_backups)."""
    cfg = cargar_config()
    carpeta_backup = ruta_cierre(cfg["backup_manual"]["carpeta_backups"])
    os.makedirs(carpeta_backup, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    nombre = os.path.basename(ruta_archivo)
    ruta_backup = os.path.join(carpeta_backup, f"{timestamp}_{nombre}")
    shutil.copy2(ruta_archivo, ruta_backup)
    print(f"  Backup creado: {ruta_backup}")
    return ruta_backup


def _validar_referencias_estructuradas(wb) -> None:
    """Mismo guardrail que pipeline.escritura_asientos._validar_referencias_estructuradas
    (ver docstring ahi) - las hojas raw (sin formulas) se saltean."""
    nombres_reales = _tablas_reales(wb)
    problemas = []
    for nombre_hoja, config_hoja in HOJAS_FILTRADAS.items():
        ws = wb[nombre_hoja]
        tabla = ws.tables[config_hoja["tabla"]]
        _, ref_fin = tabla.ref.split(":")
        fila_max = max(int(re.search(r"\d+", ref_fin).group()), _ultima_fila_real(ws))
        col_formula_idx = ws[config_hoja["primera_col_formula"] + "1"].column
        col_fin_idx = ws[re.match(r"[A-Za-z]+", ref_fin).group() + "1"].column

        for fila in ws.iter_rows(min_row=FILA_TEMPLATE, max_row=fila_max,
                                  min_col=col_formula_idx, max_col=col_fin_idx):
            for celda in fila:
                valor = celda.value
                if not isinstance(valor, str) or not valor.startswith("="):
                    continue
                for m in _TABLE_REF_RE.finditer(valor):
                    nombre_tabla = m.group(1)
                    if nombre_tabla not in nombres_reales:
                        problemas.append(
                            f"{nombre_hoja}!{celda.coordinate}: referencia a tabla "
                            f"inexistente '{nombre_tabla}' en formula {valor!r}"
                        )

    if problemas:
        raise ValueError(
            "Guardrail de referencias estructuradas: se encontraron formulas que "
            "apuntan a tablas inexistentes despues de pegar - NO se guardo el archivo "
            f"de cierre (evita corromperlo). Revisar _ajustar_filas_formula:\n"
            + "\n".join(problemas[:20])
            + (f"\n... y {len(problemas) - 20} mas" if len(problemas) > 20 else "")
        )


def pegar_todo_manual(df_acumulaciones: pd.DataFrame, df_redenciones: pd.DataFrame,
                       df_puntos_expirados: pd.DataFrame, point_types_acumulacion_non_tender: list,
                       point_types_redenciones_general: list, point_types_redenciones_subs: list,
                       df_diccionario_excl: pd.DataFrame = None) -> dict:
    """Backup + abrir + pegar las 7 hojas (+ Diccionario si se pasa) + validar +
    guardar. Punto de entrada unico, igual patron que pipeline.escritura_asientos.pegar_todo()."""
    from openpyxl import load_workbook

    ruta = ruta_cierre_mensual()
    hacer_backup(ruta)

    wb = load_workbook(ruta)  # modo normal: seguro con el archivo unificado sin pivots (~2.1 GB)
    resultado = {
        "Query Acumulaciones": pegar_query_acumulaciones(wb, df_acumulaciones),
        "Query Redenciones": pegar_query_redenciones(wb, df_redenciones),
        "Breakage": pegar_breakage(wb, df_puntos_expirados),
        "Acumulacion Non Tender": pegar_acumulacion_non_tender(
            wb, df_acumulaciones, point_types_acumulacion_non_tender
        ),
    }
    resultado.update(pegar_redenciones_filtradas(
        wb, df_redenciones, point_types_redenciones_general, point_types_redenciones_subs
    ))
    if df_diccionario_excl is not None:
        resultado["Diccionario"] = pegar_diccionario(wb, df_diccionario_excl)
    _validar_referencias_estructuradas(wb)
    wb.save(ruta)
    print(f"  Guardado OK: {ruta}")
    return resultado
