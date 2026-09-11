"""
Etapa 1 - Paso final: pega los resultados de Acumulacion Non Tender / Redenciones /
Redenciones SUBS / Redenciones Otros en Asientos Cierre Loyalty.xlsx.

IMPORTANTE (confirmado por Rosario, 2026-07-21) - el archivo NO acumula historico entre
meses. Cada carpeta de mes (duplicada a mano de la carpeta del mes anterior, ver
Bitacora seccion 15) tiene SU PROPIO Asientos Cierre Loyalty.xlsx, y cada corrida debe
dejarlo con SOLO los datos de ESE mes: se borra la base pegada a valor del mes anterior
(heredada por la copia) y se inserta la del mes nuevo. Del archivo heredado solo se
reusa la ESTRUCTURA (headers, Tabla de Excel, y el patron de formulas/valores
constantes del asiento contable) - no su contenido. Esto reemplaza el diseño anterior
(que asumia acumulacion y solo agregaba filas al final - INCORRECTO, corregido esta
sesion antes de tocar ningun archivo real).

Mecanica por hoja (_pegar_hoja):
1. Capturar el patron de formulas/valores constantes del bloque de asiento contable
   fijo DESDE LA FILA 3 (fila plantilla) - ANTES de borrar nada, porque una vez borrado
   ya no queda de donde copiarlo.
2. Escribir los datos del mes nuevo (bloque raw + calculadas, por nombre de columna)
   arrancando siempre en la fila 3 - pisando lo que hubiera ahi.
3. Para cada fila nueva, replicar el patron capturado en el paso 1 (formulas con el
   numero de fila ajustado, valores constantes tal cual - ver _ajustar_filas_formula).
4. Si el mes anterior tenia MAS filas que el mes nuevo, borrar el remanente (sino
   quedarian pegadas filas del mes anterior despues de las nuevas).
5. Ajustar el rango de la Tabla de Excel (`ref`) al tamano exacto del mes nuevo (puede
   crecer o achicarse respecto del mes anterior).

Mapeo de columnas POR NOMBRE (no por posicion) contra el header real (fila 2) de cada
hoja. OJO: el header real de Asientos usa nombres DISTINTOS a los que trae el DataFrame
en varios casos (confirmado inspeccionando el archivo real, 2026-07-21) - por eso el
mapeo de abajo es una lista explicita (header en Asientos, columna en el DataFrame), no
una simple igualdad de nombres:
  - Redenciones / Redenciones Subs / Redenciones Otros: Asientos llama "country" a lo
    que el DataFrame trae como `country_code`, "pais" a lo que trae como `country`, y
    "tipopago" a lo que trae como `payment_type`.
  - Redenciones Otros ademas: "parentchannel" en vez de `channel_condition`, y
    "Descuento por Consumo de PuntosML" en vez de `Reconocimiento de Ingresos Diferidos ML`.

Cada hoja tiene una Tabla de Excel (ListObject): Tabla1 (Generacion), Tabla2
(Redenciones), Tabla28 (Redenciones Subs), Tabla25 (Redenciones Otros). Su `ref` se
usa para saber cuantas filas tenia el mes anterior (no `ws.max_row`, que esta inflado
por filas con formato pero sin datos).

Se hace un backup con timestamp del archivo ANTES de escribir, siempre (ver
config.yaml -> rutas.carpeta_backups_asientos) - imprescindible ahora que el proceso
borra datos existentes, no solo agrega.

NOTA (2026-07-21): se encontraron 2 formulas de control en el archivo (Redenciones!AP1
y Redenciones Subs!AP1, con un rango que arranca muy por delante del final real de los
datos) - Rosario confirmo que estan bien formuladas asi (da 0 a proposito en algunos
casos). No se tocan.
"""

import os
import re
import shutil
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from pipeline.config import cargar_config, ruta_cierre


def _mapeo_redenciones(header_channel="channel_condition",
                        header_reconocimiento_ml="Reconocimiento de Ingresos Diferidos ML"):
    """2026-07-22: se sacaron las 6 entradas de columnas eliminadas de la query de
    Redenciones (dsp_transaction_id, gb_total_carrito, peso_producto, transaction_code,
    transaction_type, ratio_prd) - Rosario ya restructuro Asientos y esos headers no
    existen mas ahi. El orden de la lista no importa (se mapea por nombre, no por
    posicion) - "Entidad Legal" ahora vive al lado de "Descuento por Consumo de Puntos"
    en el archivo real, pero no hizo falta cambiar nada de esta entrada especifica."""
    return [
        ("processing_date", "processing_date"),
        ("product", "product"),
        ("country", "country_code"),
        ("partner", "partner"),
        ("points", "points"),
        ("point_type", "point_type"),
        ("points_distribuidos", "points_distribuidos"),
        ("business", "business"),
        ("pais", "country"),
        ("produto_original", "produto_original"),
        (header_channel, "channel_condition"),
        ("tipopago", "payment_type"),
        ("produto_agrupado", "produto_agrupado"),
        ("comision", "comision"),
        ("fee", "fee"),
        ("descuentos", "descuentos"),
        ("gb_basebi", "gb_basebi"),
        ("gb_basebi_2", "gb_basebi_2"),
        ("Descuento por Consumo de Puntos", "Descuento por Consumo de Puntos"),
        ("Entidad Legal", "Entidad Legal"),
        (header_reconocimiento_ml, "Reconocimiento de Ingresos Diferidos ML"),
        ("Concat Entidad Legal", "Concat Entidad Legal"),
        ("Producto", "Producto"),
    ]


MAPEO_GENERACION = [
    ("processing_date", "processing_date"),
    ("country_code", "country_code"),
    ("partner", "partner"),
    ("points", "points"),
    ("point_type", "point_type"),
    ("business", "business"),
    ("product", "product"),
    ("country", "country"),
    ("payment_type", "payment_type"),
    ("channel_condition", "channel_condition"),
    ("trip_type", "trip_type"),
    ("produto_original", "produto_original"),
    ("produto", "produto"),
    ("comision", "comision"),
    ("fee", "fee"),
    ("descuentos", "descuentos"),
    ("gb_basebi", "gb_basebi"),
    ("gb_basebi_2", "gb_basebi_2"),
    ("pct_pagado_con_puntos", "pct_pagado_con_puntos"),
    ("Revenue", "Revenue"),
    ("Puntos Valuados", "Puntos Valuados"),  # Etapa 2 pendiente - el DataFrame no la trae, queda vacia
    ("DRO", "DRO"),                          # idem
    ("DRO - up fronts", "DRO - up fronts"),  # idem
    ("DRO -Fee + Descuentos", "DRO -Fee + Descuentos"),  # idem
    ("Filtro P&L", "Filtro P&L"),
    ("Entidad Legal", "Entidad Legal"),
    ("Producto", "Producto"),
    ("Points ABS", "Points ABS"),
]

MAPEO_REDENCIONES = _mapeo_redenciones()
MAPEO_REDENCIONES_SUBS = _mapeo_redenciones()
MAPEO_REDENCIONES_OTROS = _mapeo_redenciones(
    header_channel="parentchannel",
    header_reconocimiento_ml="Descuento por Consumo de PuntosML",
)

# Breakage por pais/Entidad Legal (ver Proceso Cierre.pdf y Bitacora seccion 44): las
# columnas K:R de 'Breakage' en el archivo de cierre van a valor en Asientos!Breakage
# A:H - a diferencia de Generacion/Redenciones (que salen de Etapa 1 + columnas Etapa 2
# superpuestas), este bloque es 100% formula en el archivo de cierre (SUMIFS/XLOOKUP
# contra TablaANT[DRO]) y se lee entero de una - ver
# pipeline_backup.lectura_etapa2.leer_breakage_pais_el().
MAPEO_BREAKAGE = [
    ("country", "country_code"),
    ("Entidad Legal", "Entidad Legal"),
    ("channel_condition", "channel_condition"),
    ("produto_original", "produto_original"),
    ("Producto", "Producto"),
    ("Suma de DRO", "Suma de DRO"),
    ("% DRO", "% DRO"),
    ("Breakage por país y LE (USD)", "Breakage por país y LE (USD)"),
]

HOJAS = {
    "Generacion": {"tabla": "Tabla1", "mapeo": MAPEO_GENERACION, "primera_col_formula": "AC"},
    "Redenciones": {"tabla": "Tabla2", "mapeo": MAPEO_REDENCIONES, "primera_col_formula": "AD"},
    "Redenciones Subs": {"tabla": "Tabla28", "mapeo": MAPEO_REDENCIONES_SUBS, "primera_col_formula": "AD"},
    "Redenciones Otros": {"tabla": "Tabla25", "mapeo": MAPEO_REDENCIONES_OTROS, "primera_col_formula": "AD"},
    "Breakage": {"tabla": "Tabla3", "mapeo": MAPEO_BREAKAGE, "primera_col_formula": "I"},
}


def ruta_asientos() -> str:
    cfg = cargar_config()
    return ruta_cierre(cfg["rutas"]["asientos_cierre_loyalty"])


def hacer_backup(ruta_archivo: str) -> str:
    cfg = cargar_config()
    carpeta_backup = ruta_cierre(cfg["rutas"]["carpeta_backups_asientos"])
    os.makedirs(carpeta_backup, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    nombre = os.path.basename(ruta_archivo)
    ruta_backup = os.path.join(carpeta_backup, f"{timestamp}_{nombre}")
    shutil.copy2(ruta_archivo, ruta_backup)
    print(f"  Backup creado: {ruta_backup}")
    return ruta_backup


_CELL_REF_RE = re.compile(r"(?<![A-Za-z0-9_])(\$?)([A-Za-z]{1,3})(\$?)(\d+)")


def _ajustar_filas_formula(formula: str, fila_origen: int, fila_destino: int) -> str:
    """Ajusta las referencias de fila NO absolutas (sin $) de una formula, igual que un
    arrastre de Excel desde fila_origen hacia fila_destino. Las referencias con $ en la
    fila (ej. $B$4, vlookup!$W$3:$X$19) quedan intactas a proposito.

    2026-07-22: se agrego un negative lookbehind `(?<![A-Za-z0-9_])` antes del grupo de
    letras - bug real encontrado en la corrida real contra la carpeta TRIAL. Sin esto,
    el regex matcheaba el SUFIJO de un nombre de Tabla de Excel como si fuera una
    referencia de celda: en "Tabla2[[#This Row],[...]]" matcheaba "bla2" (col="bla",
    fila="2", patron valido de celda aunque "bla" no sea una columna real), y al
    ajustar la fila terminaba reescribiendo el nombre de la tabla completo (Tabla2 ->
    Tabla3 -> Tabla4 -> Tabla5... segun cuantas filas de arrastre). Como esas tablas
    con numeracion mas alta no existen (o son OTRAS tablas del archivo, sin relacion),
    la formula estructurada quedaba rota -> #REF! a partir de la fila donde el numero
    incrementado ya no correspondia a ninguna tabla real - coincide exactamente con lo
    reportado (las primeras 1-2 filas "por casualidad" apuntaban a una tabla que si
    existia, aunque incorrecta, y de ahi en mas #REF!). Con el lookbehind, un nombre de
    tabla completo (ej. "Tabla2", "Tabla25") nunca matchea como celda, porque ninguna
    posicion intermedia de esas letras esta libre de una letra inmediatamente antes."""
    delta = fila_destino - fila_origen

    def reemplazar(m):
        col_dollar, col, row_dollar, row = m.groups()
        if row_dollar == "$":
            return m.group(0)
        return f"{col_dollar}{col}{row_dollar}{int(row) + delta}"

    return _CELL_REF_RE.sub(reemplazar, formula)


def _set_cell(ws, row: int, column: int, value, number_format: str = None) -> None:
    """Asigna el valor SIEMPRE, incluso si es None (a diferencia de
    ws.cell(row, column, value=None), que openpyxl trata como 'no hacer nada' - ver
    su implementacion: solo asigna si value is not None. Con .value = None directo
    si se limpia la celda de verdad. Bug real encontrado 2026-07-21 al implementar
    el borrado del mes anterior: sin este helper, las celdas que debian quedar
    vacias conservaban el valor viejo.

    `number_format`, si se pasa, fija el formato de la celda en vez de dejar el que
    ya tuviera (ver FORMATOS_FIJOS - columnas donde el formato heredado de la fila
    vieja es inconsistente entre bloques de filas, ej. pct_pagado_con_puntos que a
    veces se ve como porcentaje y a veces como decimal crudo segun la fila)."""
    celda = ws.cell(row=row, column=column)
    celda.value = value
    if number_format is not None:
        celda.number_format = number_format


FORMATOS_FIJOS = {
    # columna del DataFrame -> number_format fijo, sin importar el formato heredado
    # de la celda (ver docstring de _set_cell).
    "pct_pagado_con_puntos": "0.00%",
}


def _header_a_columna(ws, fila_header: int = 2) -> dict:
    """{texto del header: indice de columna 1-based} leyendo la fila de headers real."""
    mapa = {}
    for cell in ws[fila_header]:
        if cell.value is not None:
            mapa[str(cell.value).strip()] = cell.column
    return mapa


FILA_TEMPLATE = 3  # primera fila de datos - fuente del patron de formulas/constantes
FILA_HEADER = 2


def _pegar_hoja(wb, nombre_hoja: str, df: pd.DataFrame) -> int:
    """Deja la hoja con SOLO los datos de df (borra lo que hubiera del mes anterior).
    Devuelve la cantidad de filas escritas."""
    config_hoja = HOJAS[nombre_hoja]
    ws = wb[nombre_hoja]
    tabla = ws.tables[config_hoja["tabla"]]

    col_inicio_ref, ref_fin = tabla.ref.split(":")
    fila_vieja_max = int(re.search(r"\d+", ref_fin).group())
    letra_col_fin = re.match(r"[A-Za-z]+", ref_fin).group()

    encabezados = _header_a_columna(ws, FILA_HEADER)
    primera_col_formula_idx = ws[config_hoja["primera_col_formula"] + "1"].column
    ultima_col_idx = max(encabezados.values())

    # 1. Capturar el patron de formulas/constantes del asiento contable ANTES de
    # borrar nada - una vez que se pisan los datos del mes anterior, ya no queda de
    # donde copiarlo.
    plantilla = {}
    for col_idx in range(primera_col_formula_idx, ultima_col_idx + 1):
        plantilla[col_idx] = ws.cell(row=FILA_TEMPLATE, column=col_idx).value

    df = df.reset_index(drop=True) if df is not None else pd.DataFrame()
    filas_nuevas = len(df)
    fila_nueva_max = FILA_HEADER + filas_nuevas  # FILA_HEADER + 0 si no hay filas nuevas

    # 2. Escribir los datos del mes nuevo, siempre arrancando en FILA_TEMPLATE (pisa
    # lo que hubiera del mes anterior).
    for i, fila in df.iterrows():
        fila_excel = FILA_TEMPLATE + i

        for header_asientos, columna_df in config_hoja["mapeo"]:
            col_idx = encabezados.get(header_asientos)
            if col_idx is None:
                raise KeyError(
                    f"No se encontro el header '{header_asientos}' en la hoja "
                    f"'{nombre_hoja}' de Asientos Cierre Loyalty.xlsx - revisar si "
                    "cambio el archivo antes de seguir."
                )
            valor = fila[columna_df] if columna_df in df.columns else None
            if pd.isna(valor):
                valor = None
            _set_cell(ws, fila_excel, col_idx, valor, FORMATOS_FIJOS.get(columna_df))

        # 3. Replicar la plantilla del asiento contable (formulas ajustadas / valores
        # constantes tal cual) en cada fila nueva.
        for col_idx, valor_plantilla in plantilla.items():
            if valor_plantilla is None:
                nuevo_valor = None
            elif isinstance(valor_plantilla, str) and valor_plantilla.startswith("="):
                nuevo_valor = _ajustar_filas_formula(valor_plantilla, FILA_TEMPLATE, fila_excel)
            else:
                nuevo_valor = valor_plantilla
            _set_cell(ws, fila_excel, col_idx, nuevo_valor)

    # 4. Si el mes anterior tenia MAS filas que el mes nuevo, borrar el remanente
    # (todas las columnas, de A a la ultima) para que no queden filas viejas colgando.
    if fila_vieja_max > fila_nueva_max:
        for fila_excel in range(fila_nueva_max + 1, fila_vieja_max + 1):
            for col_idx in range(1, ultima_col_idx + 1):
                _set_cell(ws, fila_excel, col_idx, None)

    # 5. Ajustar el rango de la Tabla al tamano exacto del mes nuevo.
    tabla.ref = f"{col_inicio_ref}:{letra_col_fin}{fila_nueva_max}"

    if filas_nuevas:
        print(f"  {nombre_hoja}: {filas_nuevas} fila(s) del mes (filas {FILA_TEMPLATE} a "
              f"{fila_nueva_max}) - datos del mes anterior reemplazados.")
    else:
        print(f"  {nombre_hoja}: 0 filas este mes - datos del mes anterior borrados, hoja vacia.")
    return filas_nuevas


def pegar_generacion(wb, df_acumulacion_non_tender: pd.DataFrame) -> int:
    return _pegar_hoja(wb, "Generacion", df_acumulacion_non_tender)


def pegar_redenciones(wb, df_redenciones: pd.DataFrame, df_redenciones_subs: pd.DataFrame,
                       df_redenciones_otros: pd.DataFrame) -> dict:
    return {
        "Redenciones": _pegar_hoja(wb, "Redenciones", df_redenciones),
        "Redenciones Subs": _pegar_hoja(wb, "Redenciones Subs", df_redenciones_subs),
        "Redenciones Otros": _pegar_hoja(wb, "Redenciones Otros", df_redenciones_otros),
    }


def pegar_breakage(wb, df_breakage: pd.DataFrame) -> int:
    return _pegar_hoja(wb, "Breakage", df_breakage)


_TABLE_REF_RE = re.compile(r"(?<![A-Za-z0-9_.])([A-Za-z_][A-Za-z0-9_.]*)\[")


def _tablas_reales(wb) -> set:
    return {tabla.name for ws in wb.worksheets for tabla in ws.tables.values()}


def _validar_referencias_estructuradas(wb) -> None:
    """Guardrail: escanea el bloque de formulas de cada hoja pegada buscando
    referencias estructuradas (`NombreTabla[...]`) a tablas que NO existen en el
    workbook. Frena y NO guarda si encuentra alguna.

    Por que existe (2026-07-22/23): un bug en `_ajustar_filas_formula` (regex sin
    limite de palabra, ya corregido - ver docstring de esa funcion) mangleaba el
    nombre de una Tabla real al arrastrar formulas (Tabla2 -> Tabla3 -> Tabla4 ->
    ...), generando miles de referencias a tablas inexistentes -> #REF! en el
    archivo. Ese archivo corrupto, al abrirse en Excel, disparaba una reparacion
    global del paquete OOXML que ademas invalidaba TODOS los vinculos externos del
    archivo (las celdas que apuntaban a otro libro pasaban a mostrar
    "[Vinculo Externo Recuperado]") - sintoma reportado por Rosario, sin relacion
    aparente con el bug real de fondo. El bug de raiz ya esta corregido, pero este
    chequeo es la defensa para que CUALQUIER bug futuro de arrastre de formulas
    (no solo este) se detecte ANTES de sobreescribir el archivo real, en vez de
    descubrirse recien cuando alguien lo abre en Excel."""
    nombres_reales = _tablas_reales(wb)
    problemas = []
    for nombre_hoja, config_hoja in HOJAS.items():
        ws = wb[nombre_hoja]
        tabla = ws.tables[config_hoja["tabla"]]
        _, ref_fin = tabla.ref.split(":")
        fila_max = int(re.search(r"\d+", ref_fin).group())
        col_formula_idx = ws[config_hoja["primera_col_formula"] + "1"].column
        col_fin_idx = ws[re.match(r"[A-Za-z]+", ref_fin).group() + "1"].column

        for fila in ws.iter_rows(min_row=FILA_TEMPLATE, max_row=fila_max,
                                  min_col=col_formula_idx, max_col=col_fin_idx):
            for celda in fila:
                valor = celda.value
                if not isinstance(valor, str) or not valor.startswith("="):
                    continue
                for m in _TABLE_REF_RE.finditer(valor):
                    nombre_tabla = m.group(1)
                    if nombre_tabla not in nombres_reales:
                        problemas.append(
                            f"{nombre_hoja}!{celda.coordinate}: referencia a tabla "
                            f"inexistente '{nombre_tabla}' en formula {valor!r}"
                        )

    if problemas:
        raise ValueError(
            "Guardrail de referencias estructuradas: se encontraron formulas que "
            "apuntan a tablas inexistentes despues de pegar - NO se guardo el archivo "
            f"(evita corromper {os.path.basename(ruta_asientos())}). Revisar "
            "_ajustar_filas_formula:\n" + "\n".join(problemas[:20])
            + (f"\n... y {len(problemas) - 20} mas" if len(problemas) > 20 else "")
        )


def pegar_accounting(wb, valores: list) -> int:
    """Pisa la hoja 'Accounting' de Asientos Cierre Loyalty.xlsx con los valores YA
    RECALCULADOS de la solapa homonima del archivo de cierre - copia celda a celda,
    a valor (reemplaza el mecanismo de vinculo externo vivo que tenia antes). Ver
    Bitacora seccion 31 - decision de Rosario: Accounting es una copia exacta de la
    del archivo de cierre.

    IMPORTANTE: la fuente (archivo de cierre) tiene mas columnas que el destino
    (confirmado: A:L en el pesado vs A:E en Asientos - las columnas de mas en la
    fuente son ayudas/referencias tipo "partner"/"point_type", no forman parte del
    asiento). Por eso NO se copia el ancho de la fuente - se respeta el ancho
    EXISTENTE de Asientos!Accounting (`ws.max_column`, tomado ANTES de escribir) y
    se recorta cualquier fila de la fuente que tenga mas columnas que eso.

    `valores` es una lista de filas (cada una una lista de valores), tal como la
    devuelve pipeline_backup.lectura_etapa2.leer_accounting()."""
    ws = wb["Accounting"]

    # La hoja tiene celdas combinadas (texto explicativo en algunas filas) - openpyxl
    # no permite escribir en una MergedCell (solo la celda ancla es escribible).
    # Como pasa a ser una grilla de valores plana, se desarman los combinados antes
    # de pegar.
    for rango in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rango))

    fila_vieja_max = ws.max_row
    columnas_destino = ws.max_column

    filas_nuevas = len(valores)

    for i, fila in enumerate(valores, start=1):
        for j in range(1, columnas_destino + 1):
            valor = fila[j - 1] if j - 1 < len(fila) else None
            _set_cell(ws, i, j, valor)

    fila_max = max(fila_vieja_max, filas_nuevas)
    for fila_excel in range(filas_nuevas + 1, fila_max + 1):
        for col_idx in range(1, columnas_destino + 1):
            _set_cell(ws, fila_excel, col_idx, None)

    return filas_nuevas


def pegar_todo(df_generacion: pd.DataFrame, df_redenciones: pd.DataFrame,
               df_redenciones_subs: pd.DataFrame, df_redenciones_otros: pd.DataFrame,
               valores_accounting: list = None, df_breakage: pd.DataFrame = None) -> dict:
    """Backup + abrir + pegar las 4 hojas (+ Breakage y Accounting si se pasan) +
    validar + guardar. Punto de entrada unico.

    `valores_accounting` (opcional, ver Bitacora seccion 31) y `df_breakage`
    (opcional, ver Bitacora seccion 44): dependen de la solapa homonima del archivo
    de cierre YA RECALCULADA (Etapa 2) - las usa escribir_asientos.py. Si se omiten
    (caso de run_cierre.py / run_cierre_backup_manual.py, que no llegan a tener
    Etapa 2 resuelta en la misma corrida), esas hojas no se tocan."""
    ruta = ruta_asientos()
    hacer_backup(ruta)

    wb = load_workbook(ruta)  # modo normal (no read_only): necesario para escribir y tocar Tablas
    resultado = {"Generacion": pegar_generacion(wb, df_generacion)}
    resultado.update(pegar_redenciones(wb, df_redenciones, df_redenciones_subs, df_redenciones_otros))
    if df_breakage is not None:
        resultado["Breakage"] = pegar_breakage(wb, df_breakage)
    if valores_accounting is not None:
        resultado["Accounting"] = pegar_accounting(wb, valores_accounting)
    _validar_referencias_estructuradas(wb)
    wb.save(ruta)
    print(f"  Guardado OK: {ruta}")
    return resultado
